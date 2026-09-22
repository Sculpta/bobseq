#!/usr/bin/env python3
"""Supplementary Table S1: one row per plotted dataset, from the pipeline output.

Inputs
  --results-dir  folder written by speciesmix_bobcode_metrics.py (bobcode_metrics.tsv and
                 <run>/<barcode>/metrics.json)
  --index        barcode_accuracy_index.tsv from make_accuracy_figure.py (plot numbers)
  --run-rules    bobcode_run_rules.tsv (codes, construct, dates, conditions = sample_notes,
                 per-barcode flags = barcode_flags, exclusions)
Outputs (--out-dir)
  Supplementary_Table_S1.xlsx   sheets: S1 Datasets, Column definitions, Excluded, Method
  Supplementary_Table_S1.tsv    the S1 Datasets sheet as plain text

Requires openpyxl for the .xlsx (the .tsv is written without it).
"""
import argparse
import csv
import json
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
ORDER = ("polydT", "splint", "5'TSO", "3'TSO")      # figure panel order
PANEL = {"polydT": "poly(dT) priming bobcode", "splint": "splint-mediated ligation of bobcode to cDNA",
         "5'TSO": "TSO bobcode (5'DBCO)", "3'TSO": "TSO bobcode (3'DBCO)"}
CONSTRUCT = {
    "bob-polydt":   "bob-c primer + 6-nt code + polyT (RT primer)",
    "bob-k-polydt": "TruSeq R1 + 12-nt code + 16N UMI + polyT (RT primer)",
    "bob-c-splint": "bob-c primer + 4-nt code + BOB_HR, splint-ligated to the RT primer",
    "bob-a":        "TruSeq R1 + 8-nt code, adaptor-1 ligation (BOB-A)",
    "nextera-tso":  "Nextera-R1 TSO + 7-nt code + rGrGrG",
    "smarter-4bp":  "SMARTer TSO primer + 4-nt code",
    "tso-bob":      "SMARTer TSO backbone + 7-nt code + rGrGrG",
    "bob-v1":       "TruSeq R1 handle + 7-nt code + UMI/spacer + rGrGrG (TSO)",
}
COLLISION = {"tso-bob": "Human code AGTACAT equals the slot of the standard SMARTer TSO; unbarcoded TSO would score as human",
             "smarter-4bp": "Mouse code TACA equals the slot of the standard SMARTer TSO; unbarcoded TSO would score as mouse"}

COLS = [  # (header, key, width, number format, definition)
 ("Chemistry", "chem", 22, None, "Panel of the barcode accuracy figure."),
 ("Plot #", "plot", 7, "0", "Number on the figure's x axis; numbered from 1 within each chemistry, ranked by accuracy. A dataset is identified by Chemistry + Plot #."),
 ("Run", "run", 34, None, "Sequencing run folder."),
 ("Sequencing date", "date", 11, None, "Date the MinKNOW run started."),
 ("ONT barcode", "bc", 8, None, "Oxford Nanopore native barcode of the library."),
 ("Construct", "construct", 44, None, "Where the species code sits (see bobcode_general_rules.md §6)."),
 ("Human code", "hcode", 13, None, "Species code assigned to human."),
 ("Mouse code", "mcode", 13, None, "Species code assigned to mouse."),
 ("Condition", "cond", 36, None, "Sample description from the run rules (sample_notes); 'condition variant' where none was recorded."),
 ("Total reads", "total", 10, "#,##0", "Reads in the raw FASTQ(s) used."),
 ("Mapped reads", "mapped", 10, "#,##0", "Reads with at least one alignment (STARlong, combined human + mouse genome)."),
 ("Human reads", "human", 10, "#,##0", "Mapped reads aligned only to human."),
 ("Mouse reads", "mouse", 10, "#,##0", "Mapped reads aligned only to mouse."),
 ("Ambiguous reads", "amb", 10, "#,##0", "Mapped reads with alignments on both species (not scored)."),
 ("mRNA reads", "mrna", 10, "#,##0", "Mapped reads on exons of protein-coding genes (ribosomal-protein genes excluded), any species."),
 ("rRNA reads", "rrna", 10, "#,##0", "Mapped reads on rDNA loci or rRNA genes."),
 ("MT reads", "mt", 9, "#,##0", "Mapped reads on the mitochondrial genome."),
 ("Other reads", "other", 10, "#,##0", "All other mapped reads (lncRNA and other non-coding, ribosomal-protein mRNA, intronic, intergenic)."),
 ("Reads scored (n)", "n", 10, "#,##0", "Species-unique mRNA reads, not low-complexity, with exactly one barcode unit and an exact code."),
 ("Correct", "correct", 10, "#,##0", "Scored reads whose code matches their aligned species."),
 ("Accuracy (%)", "acc", 9, "0.00", "Correct / Reads scored; the value plotted."),
 ("95% CI low (%)", "lo", 9, "0.00", "Wilson 95% interval on the accuracy (read-level)."),
 ("95% CI high (%)", "hi", 9, "0.00", "Upper bound of the same interval."),
 ("Human-mRNA accuracy (%)", "hacc", 10, "0.00", "Accuracy on scored reads aligned to human."),
 ("Mouse-mRNA accuracy (%)", "macc", 10, "0.00", "Accuracy on scored reads aligned to mouse."),
 ("Notes", "notes", 90, None, "Flags, direction notes (a code whose declared species is not the majority of its reads), pooling, subsampling and known code collisions."),
]


def parse_bc_map(text):
    return {k.strip(): v.strip() for k, v in (t.split(":", 1) for t in (text or "").split("|") if ":" in t)}


def build_rows(results_dir, index, run_rules):
    rules = {(r["chemistry"], r["run_folder"].split("/")[-1]): r
             for r in csv.DictReader(open(run_rules), delimiter="\t")}
    rows = []
    for xi in csv.DictReader(open(index), delimiter="\t"):
        chem, run, bc = xi["chemistry"], xi["run"], xi["barcode"]
        rr = rules[(chem, run)]
        res = json.load(open(os.path.join(results_dir, run, bc, "metrics.json")))
        mt, ct = res["metrics"], res["counts"]
        notes = []
        flag = parse_bc_map(rr.get("barcode_flags")).get(bc)
        if flag:
            notes.append(flag)
        if mt["status"] != "ok":
            notes.append(mt["status"].replace("note: ", "Direction note: "))
        if ";" in rr["raw_fastq_input"]:
            notes.append(f"{rr['raw_fastq_input'].count(';') + 1} sequencing runs of the same libraries pooled")
        if (rr.get("max_reads") or "").strip().isdigit():
            notes.append(f"Only the first {int(rr['max_reads']):,} reads are used")
        if rr["construct"] in COLLISION:
            notes.append(COLLISION[rr["construct"]])
        rows.append(dict(
            chem=PANEL[chem], plot=int(xi["n"]), run=run, date=rr["seq_date"], bc=bc,
            construct=CONSTRUCT.get(rr["construct"], rr["construct"]),
            hcode=rr["human_code"], mcode=rr["mouse_code"],
            cond=parse_bc_map(rr.get("sample_notes")).get(bc) or "condition variant",
            total=mt["reads_total"], mapped=mt["reads_mapped"],
            human=ct.get("species_human", 0), mouse=ct.get("species_mouse", 0), amb=ct.get("species_ambiguous", 0),
            mrna=mt["mRNA_reads"], rrna=mt["rRNA_reads"], mt=mt["MT_reads"], other=mt["other_reads"],
            n=mt["total_mRNA_scored"], correct=res["correct"]["total_mRNA"],
            acc=mt["total_mRNA_accuracy_pct"], lo=mt["total_ci95_lo_pct"], hi=mt["total_ci95_hi_pct"],
            hacc=mt["human_mRNA_accuracy_pct"], macc=mt["mouse_mRNA_accuracy_pct"],
            notes="; ".join(notes) or "none", _chem=chem))
    rows.sort(key=lambda x: (ORDER.index(x["_chem"]), x["plot"]))
    for c in ORDER:     # numbering restarts at 1 in each chemistry, without gaps
        p = [x["plot"] for x in rows if x["_chem"] == c]
        assert p == list(range(1, len(p) + 1)), c
    return rows, rules


def excluded(rules):
    out = []
    for (chem, run), r in sorted(rules.items(), key=lambda kv: (ORDER.index(kv[0][0]), kv[1]["seq_date"])):
        if r["barcodes_excluded (reason)"].strip() not in ("", "none"):
            out.append((PANEL[chem], run, r["barcodes_excluded (reason)"].strip()))
    return out


def write_tsv(rows, path):
    with open(path, "w", newline="") as f:
        w = csv.writer(f, delimiter="\t")
        w.writerow([c[0] for c in COLS])
        for x in rows:
            w.writerow(["" if x[c[1]] is None else x[c[1]] for c in COLS])


def write_xlsx(rows, excl, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    FILL = {"polydT": "FFF3E0", "splint": "EEF2FA", "5'TSO": "EAF6EC", "3'TSO": "F7ECF6"}
    wb = Workbook()
    ws = wb.active; ws.title = "S1 Datasets"
    for j, c in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=j, value=c[0])
        cell.font = Font(bold=True); cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = c[2]
    for i, x in enumerate(rows, 2):
        for j, c in enumerate(COLS, 1):
            cell = ws.cell(row=i, column=j, value=x[c[1]])
            if c[3]: cell.number_format = c[3]
        ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=FILL[x["_chem"]])
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(rows) + 1}"

    def simple(title, header, data, widths):
        sh = wb.create_sheet(title)
        sh.append(header)
        for cell in sh[1]: cell.font = Font(bold=True)
        for d in data: sh.append(list(d))
        for j, wdt in enumerate(widths, 1):
            sh.column_dimensions[get_column_letter(j)].width = wdt
            for cell in sh[get_column_letter(j)]: cell.alignment = Alignment(wrap_text=True, vertical="top")
    simple("Column definitions", ["Column", "Definition"], [(c[0], c[4]) for c in COLS], (26, 120))
    simple("Excluded", ["Chemistry", "Run", "Excluded barcodes and reason"], excl, (30, 40, 120))
    simple("Method", ["Step", "Rule"], [
        ("Pipeline", "speciesmix_bobcode_metrics.py with bobcode_general_rules.md and bobcode_run_rules.tsv (same rules for every dataset)."),
        ("Input", "Raw ONT reads; ONT adapter and native barcode trimmed."),
        ("Alignment", "STARlong 2.7.11b, dense combined human (GRCh38) + mouse (GRCm39) index, Ensembl 113 annotation."),
        ("Species", "Human or mouse only if every alignment is on that species; reads on both are ambiguous and not scored."),
        ("mRNA", "Aligned blocks overlapping an exon of a protein-coding gene (ribosomal-protein genes excluded); only these are scored."),
        ("Barcode", "Exact 12-nt anchor + exact code, exactly one barcode unit per read; only the sample's two codes count."),
        ("Other filters", "Low-complexity alignments removed (dinucleotide entropy < 0.65)."),
        ("Accuracy", "Correct / scored, 95% Wilson interval; datasets with < 100 scored reads excluded."),
    ], (16, 120))
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--results-dir", required=True)
    ap.add_argument("--index", required=True, help="barcode_accuracy_index.tsv from make_accuracy_figure.py")
    ap.add_argument("--run-rules", default=os.path.join(HERE, "bobcode_run_rules.tsv"))
    ap.add_argument("--out-dir", default=".")
    a = ap.parse_args()
    rows, rules = build_rows(a.results_dir, a.index, a.run_rules)
    os.makedirs(a.out_dir, exist_ok=True)
    write_tsv(rows, os.path.join(a.out_dir, "Supplementary_Table_S1.tsv"))
    try:
        write_xlsx(rows, excluded(rules), os.path.join(a.out_dir, "Supplementary_Table_S1.xlsx"))
        xl = "xlsx + tsv"
    except ImportError:
        xl = "tsv only (install openpyxl for the .xlsx)"
    print(f"S1: {len(rows)} datasets written ({xl})")


if __name__ == "__main__":
    main()
