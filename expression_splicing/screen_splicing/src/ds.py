"""Stage ds: differential splicing, five contrasts (common.CONTRASTS), on LSV-junction FRAGMENT counts (S6).

Per contrast, on its wells (3 controls incl. Hek_control_2, 3 wells per treated condition):
  prefilter (condition-blind): LSV side mean total >= MIN_MEAN_TOTAL fragments and > 0 in every well; PSI variance
                               across the wells >= MIN_VAR
  test:  pairwise -> Welch t-test on per-well PSI (treated vs control)
         dose     -> OLS of per-well PSI on dose rank 0/1/2, t-test on the slope (7 residual df)
         BH within contrast
  effect: dPSI = PSI(treated / top dose) - PSI(control) from pooled counts; dPSI on UMI-deduplicated molecules alongside
  hit:   padj < MAX_PADJ and dPSI >= HIT_DPSI  -- POSITIVE direction only (each LSV contributes a +/- pair whose members
         sum to 1; the member that goes up is the event's representative)
  top 10: hits ranked by dPSI (padj < MAX_PADJ is a gate, not a ranking), one row per LSV and per junction; when a contrast has
         fewer than 10 hits the list is padded with the largest-dPSI non-hits (flagged hit = 0)
Writes results/ds_<contrast>.tsv (every tested row), ds_hits.tsv, ds_top10.tsv, ds_summary.tsv, ds_top10.md, ds_top10.xlsx
and the `ds` section of results/screen_splicing_stats.txt."""
import os, sys, warnings
import numpy as np
from scipy import stats
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *

def slope_test(P, dose):
    """Vectorised OLS of each row of P (rows x wells) on dose; returns slope and the two-sided p of the slope t-test."""
    x = dose - dose.mean(); sxx = (x ** 2).sum(); Pc = P - P.mean(1, keepdims=True); b = (Pc * x).sum(1) / sxx
    resid = Pc - np.outer(b, x); n = P.shape[1]; s2 = (resid ** 2).sum(1) / (n - 2); se = np.sqrt(s2 / sxx)
    t = np.where(se > 0, b / np.where(se > 0, se, 1), 0.0); p = 2 * stats.t.sf(np.abs(t), n - 2); p[se == 0] = 1.0
    return b, p

def welch_test(P, ctl, tr):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", RuntimeWarning)                        # scipy warns on (near-)constant rows; they get p = 1 below
        p = stats.ttest_ind(P[:, tr], P[:, ctl], axis=1, equal_var=False).pvalue
    return np.where(np.isfinite(p), p, 1.0)                                    # both groups constant -> no evidence

def run(tag, spec, st):
    W = contrast_wells(tag); accs = [w["acc_number"] for w in W]; names = [w["sample_name"] for w in W]
    cond = np.array([w["condition"] for w in W]); groups = spec["groups"]; levels = list(groups)
    dose = np.array([groups[c] for c in cond], float); ctl = dose == 0; top = dose == max(groups.values())
    rn, K, T = load_counts(COUNTS_FRAG, accs); rn_m, K_m, T_m = load_counts(COUNTS_MOL, accs); mol_idx = {r: i for i, r in enumerate(rn_m)}
    ok = (T.mean(1) >= MIN_MEAN_TOTAL) & (T > 0).all(1)
    P = np.where(T > 0, K / np.where(T > 0, T, 1), np.nan)
    var = np.zeros(len(rn)); var[ok] = P[ok].var(1)                              # depth-passing rows have no NaN
    keep = ok & (var >= MIN_VAR); idx = np.flatnonzero(keep)
    pooled = lambda m: K[idx][:, m].sum(1) / T[idx][:, m].sum(1)
    psi_ctl, psi_top = pooled(ctl), pooled(top); dpsi = np.round(psi_top - psi_ctl, 3)
    if spec["kind"] == "dose":
        mid = dose == 1; psi_mid = pooled(mid); slope, p = slope_test(P[idx], dose)
        consistent = ((psi_ctl <= psi_mid) & (psi_mid <= psi_top)) | ((psi_ctl >= psi_mid) & (psi_mid >= psi_top))   # pooled PSIs monotone in dose
    else:
        psi_mid = np.full(len(idx), np.nan); slope = np.full(len(idx), np.nan); p = welch_test(P[idx], ctl, top)
        Pc, Pt = P[idx][:, ctl], P[idx][:, top]
        consistent = np.where(dpsi > 0, Pt.min(1) > Pc.max(1), Pt.max(1) < Pc.min(1))                                 # every treated well beyond every control well
    padj = bh(p); hit = (padj < MAX_PADJ) & (dpsi >= HIT_DPSI)
    rows = []
    for k, i in enumerate(idx):
        d = parse_row(rn[i]); mi = mol_idx.get(rn[i]); mol = ""
        if mi is not None and T_m[mi, ctl].sum() > 0 and T_m[mi, top].sum() > 0:
            mol = round(K_m[mi, top].sum() / T_m[mi, top].sum() - K_m[mi, ctl].sum() / T_m[mi, ctl].sum(), 3)
        rows.append({"contrast": tag, "rank": "", "junction_id": rn[i], **d, "mean_lsv_total": round(float(T[i].mean()), 1),
                     "psi_ctl": round(float(psi_ctl[k]), 3), "psi_mid": round(float(psi_mid[k]), 3) if spec["kind"] == "dose" else "",
                     "psi_treated": round(float(psi_top[k]), 3), "dpsi": float(dpsi[k]), "dpsi_dedup": mol,
                     "slope": round(float(slope[k]), 4) if spec["kind"] == "dose" else "", "test": "ols_slope" if spec["kind"] == "dose" else "welch",
                     "p": float(p[k]), "padj": float(padj[k]), "consistent": int(consistent[k]), "hit": int(hit[k]),
                     **{f"psi_{n}": round(float(P[i, w]), 3) for w, n in enumerate(names)}, **{f"n_{n}": int(T[i, w]) for w, n in enumerate(names)}})
    # top 10: positive dPSI only, hits first, ranked by dPSI (the padj cut is a gate, not a ranking); one row per LSV and one
    # per junction (a junction sits in both its source and its target LSV -- the better-ranked context is kept)
    pos = sorted([r for r in rows if r["dpsi"] > 0], key=lambda r: (-r["hit"], -r["dpsi"], r["padj"]))
    seen, top10 = set(), []
    for r in pos:
        if r["lsv"] in seen or r["junction"] in seen: continue
        seen.update((r["lsv"], r["junction"])); r["rank"] = len(top10) + 1; top10.append(r)
        if len(top10) == N_TOP: break
    rows.sort(key=lambda r: (r["rank"] == "", r["rank"] or 0, -r["hit"], -r["dpsi"], r["padj"]))
    cols = ["contrast", "rank", "junction_id", "gene", "lsv", "side", "anchor", "junction", "chrom", "start", "end", "strand", "mean_lsv_total", "psi_ctl", "psi_mid",
            "psi_treated", "dpsi", "dpsi_dedup", "slope", "test", "p", "padj", "consistent", "hit"] + [f"psi_{n}" for n in names] + [f"n_{n}" for n in names]
    write_tsv(os.path.join(RESULTS, f"ds_{tag}.tsv"), rows, cols)
    hits = [r for r in rows if r["hit"]]
    summ = {"contrast": tag, "title": spec["title"], "test": rows[0]["test"], "wells": len(W), "conditions": " / ".join(LAB[c] for c in levels),
            "rows_prefilter_depth": int(ok.sum()), "rows_tested": len(rows), "lsv_sides_tested": len({r["lsv"] for r in rows}),
            "padj_lt_0.1_any_direction": int((padj < MAX_PADJ).sum()), "hits_pos": len(hits), "hits_pos_lsv": len({r["lsv"] for r in hits}),
            "hits_pos_genes": len({r["gene"] for r in hits}), "top10_hits": sum(r["hit"] for r in top10),
            "hit_genes": ", ".join(sorted({r["gene"] for r in hits}))}
    st.add(f"{tag} ({spec['title']}; {rows[0]['test']}): wells {', '.join(names)}")
    st.add(f"{tag}: depth prefilter {ok.sum():,} rows; tested {len(rows):,} rows / {summ['lsv_sides_tested']:,} LSV sides; padj<{MAX_PADJ} (any direction) {summ['padj_lt_0.1_any_direction']}; "
           f"HITS (padj<{MAX_PADJ}, dPSI>={HIT_DPSI}, positive) {len(hits)} rows / {summ['hits_pos_lsv']} LSVs / {summ['hits_pos_genes']} genes: {summ['hit_genes'] or '-'}")
    st.add(f"{tag}: top10 = " + ", ".join("{}({:+.2f},q={:.2g}{})".format(r["gene"], r["dpsi"], r["padj"], "" if r["hit"] else ",ns") for r in top10))
    print(f"{tag:<11} tested {len(rows):>5}  padj<0.1 {summ['padj_lt_0.1_any_direction']:>4}  hits+ {len(hits):>3} rows / {summ['hits_pos_lsv']:>3} LSVs  top10 hits {summ['top10_hits']:>2}  "
          f"top: {', '.join(r['gene'] for r in top10[:5])}")
    return rows, hits, top10, summ, cols

def tables(top_all, cols):
    md = ["# Differential splicing — top 10 junctions per contrast (positive ΔPSI)", "",
          f"Hit = padj < {MAX_PADJ} and ΔPSI ≥ {HIT_DPSI}, positive direction. Hits ranked by ΔPSI (the padj cut is a gate, not a ranking); one row per LSV "
          "and per junction; the largest-ΔPSI non-hits pad the list when a contrast has fewer than 10 hits. PSI from pooled fragment counts; ΔPSI = treated (top dose) − control.", ""]
    for tag, spec in CONTRASTS.items():
        sub = [r for r in top_all if r["contrast"] == tag]
        md += [f"## {spec['title']} — {sum(r['hit'] for r in sub)} of the 10 are hits", "",
               "| rank | gene | LSV | junction | PSI control | PSI treated | ΔPSI | padj | hit |", "|---|---|---|---|---|---|---|---|---|"]
        md += [f"| {r['rank']} | {r['gene']} | {r['lsv']} | {r['junction']} | {r['psi_ctl']:.2f} | {r['psi_treated']:.2f} | {r['dpsi']:+.2f} | {r['padj']:.2g} | {'yes' if r['hit'] else 'no'} |" for r in sub]
        md.append("")
    open(os.path.join(RESULTS, "ds_top10.md"), "w").write("\n".join(md))
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    xc = [("rank", "rank", 6), ("gene", "gene", 14), ("lsv", "LSV", 28), ("junction", "junction (chr:start-end:strand)", 28), ("psi_ctl", "PSI control", 11),
          ("psi_mid", "PSI mid dose", 11), ("psi_treated", "PSI treated / top dose", 12), ("dpsi", "ΔPSI", 8), ("dpsi_dedup", "ΔPSI (UMI molecules)", 11),
          ("p", "p", 10), ("padj", "padj", 10), ("consistent", "consistent / ordered", 11), ("hit", f"hit (padj<{MAX_PADJ} & ΔPSI≥{HIT_DPSI})", 14), ("mean_lsv_total", "LSV fragments (mean/well)", 12)]
    wb = openpyxl.Workbook(); wb.remove(wb.active); head = Font(bold=True); fill = PatternFill("solid", fgColor="DDDDDD"); hitfill = PatternFill("solid", fgColor="FDE9D9")
    ws = wb.create_sheet("README"); ws.column_dimensions["A"].width = 150
    for line in ["Differential splicing, BOBseq 24-plex (HEK, 2x150 PE): top 10 junctions per contrast, positive ΔPSI only",
                 "Pairwise contrasts: Welch t-test on per-well PSI (3 vs 3 wells; controls = Hek_control_1/2/3). Dose contrasts: OLS slope of per-well PSI on dose rank 0/1/2 (9 wells, 7 df).",
                 "BH within contrast. PSI = junction fragments / LSV-side fragments (pooled over the wells of a condition); ΔPSI = treated (top dose) − control.",
                 f"Hit = padj < {MAX_PADJ} and ΔPSI ≥ {HIT_DPSI}. Ranking: hits by ΔPSI (padj is a gate, not a ranking); one row per LSV and per junction; the largest-ΔPSI non-hits (unshaded) pad the list to 10.",
                 f"Prefilter (condition-blind): LSV side ≥ {MIN_MEAN_TOTAL} fragments on average and > 0 in every well of the contrast; PSI variance across the wells ≥ {MIN_VAR}.",
                 "Source: results/ds_top10.tsv; full tables results/ds_<contrast>.tsv."]: ws.append([line])
    for tag, spec in CONTRASTS.items():
        ws = wb.create_sheet(tag); sub = [r for r in top_all if r["contrast"] == tag]
        ws.append([f"{spec['title']} — {sub[0]['test'] if sub else ''}; {sum(r['hit'] for r in sub)} of 10 are hits"]); ws["A1"].font = Font(italic=True); ws.append([])
        ws.append([c[1] for c in xc])
        for c in ws[3]: c.font = head; c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in sub:
            ws.append([r[k] if r[k] != "" else None for k, _, _ in xc])
            if r["hit"]:
                for c in ws[ws.max_row]: c.fill = hitfill
        for i, (_, _, w) in enumerate(xc, 1): ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A4"
    wb.save(os.path.join(RESULTS, "ds_top10.xlsx"))

def main():
    os.makedirs(RESULTS, exist_ok=True); st = Stats("ds")
    st.add(f"prefilter: LSV mean total >= {MIN_MEAN_TOTAL} fragments & > 0 in all wells; PSI var >= {MIN_VAR}. hit: padj < {MAX_PADJ} & dPSI >= {HIT_DPSI} (positive only). unit: fragments (per-read build).")
    hits_all, top_all, summ_all = [], [], []
    for tag, spec in CONTRASTS.items():
        rows, hits, top10, summ, cols = run(tag, spec, st); hits_all += hits; top_all += top10; summ_all.append(summ)
    write_tsv(os.path.join(RESULTS, "ds_hits.tsv"), hits_all, cols[:24]); write_tsv(os.path.join(RESULTS, "ds_top10.tsv"), top_all, cols[:24])
    write_tsv(os.path.join(RESULTS, "ds_summary.tsv"), summ_all, list(summ_all[0].keys())); tables(top_all, cols); st.write()

if __name__ == "__main__":
    main()
